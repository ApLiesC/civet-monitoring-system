
"""Generate an Excel report from civet monitoring pipeline outputs."""

import csv
from collections import defaultdict
from datetime import datetime, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties


class BehaviorReportGenerator:
    def __init__(self, output_dir):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.report_path = self.output_dir / "civet_monitoring_report.xlsx"

        # Thin gray borders for table cells.
        self.thin_border = Border(
            left=Side(style="thin", color="BFBFBF"),
            right=Side(style="thin", color="BFBFBF"),
            top=Side(style="thin", color="BFBFBF"),
            bottom=Side(style="thin", color="BFBFBF"),
        )

    @staticmethod
    def read_csv(path):
        path = Path(path)

        if not path.exists():
            return []

        with open(path, "r", newline="", encoding="utf-8-sig") as file:
            return list(csv.DictReader(file))

    @staticmethod
    def parse_duration(value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def format_time(value):
        if value is None:
            return ""

        if isinstance(value, (datetime, time)):
            return value.strftime("%H:%M:%S")

        return str(value)

    @staticmethod
    def format_elapsed(seconds):
        seconds = max(0, int(seconds))
        hours, remainder = divmod(seconds, 3600)
        minutes, seconds = divmod(remainder, 60)

        return f"{hours:02d}:{minutes:02d}:{seconds:02d}"

    def apply_borders(self, sheet, cell_range):
        """Apply visible borders to every cell in a range."""
        for row in sheet[cell_range]:
            for cell in row:
                cell.border = self.thin_border

    @staticmethod
    def style_header(cell, fill="1F4E78"):
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(color="FFFFFF", bold=True, size=9)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    @staticmethod
    def style_subheader(cell, fill):
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    @staticmethod
    def style_title(sheet, cell_range, title):
        sheet.merge_cells(cell_range)

        cell = sheet[cell_range.split(":")[0]]
        cell.value = title
        cell.font = Font(size=15, bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        sheet.row_dimensions[1].height = 28

    @staticmethod
    def configure_print(sheet, print_area, header_rows=None):
        """Configure a worksheet to print on one A4 portrait page."""

        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(
            fitToPage=True,
            autoPageBreaks=False,
        )

        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.orientation = "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1

        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.35
        sheet.page_margins.bottom = 0.35
        sheet.page_margins.header = 0.15
        sheet.page_margins.footer = 0.15

        sheet.print_options.horizontalCentered = True
        sheet.print_options.verticalCentered = False

        sheet.print_area = print_area

        if header_rows:
            sheet.print_title_rows = header_rows

        sheet.oddFooter.center.text = "Civet Monitoring Report"
        sheet.oddFooter.right.text = "Page &P of &N"

        # Show normal Excel worksheet gridlines.
        sheet.sheet_view.showGridLines = True

    def add_activity_sheet(self, workbook, activity_records):
        sheet = workbook.active
        sheet.title = "Activity Summary"

        self.style_title(
            sheet,
            "A1:G1",
            "Civet Activity Summary",
        )

        headers = [
            "No.",
            "Inactive Start",
            "Inactive Stop",
            "Active Start",
            "Active Stop",
            "Not Visible Start",
            "Not Visible Stop",
        ]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=column, value=header)

            if column == 1:
                self.style_header(cell)
            elif column in (2, 3):
                self.style_subheader(cell, "70AD47")
            elif column in (4, 5):
                self.style_subheader(cell, "ED7D31")
            else:
                self.style_subheader(cell, "808080")

        sheet.row_dimensions[3].height = 30

        intervals = []

        for record in activity_records:
            state = record.get("state")

            if state not in {"active", "inactive", "not_visible"}:
                continue

            try:
                elapsed_start = float(record.get("start_time", 0))
            except (TypeError, ValueError):
                elapsed_start = float("inf")

            intervals.append({
                "state": state,
                "start": self.format_time(
                    record.get("cctv_start_time")
                ),
                "end": self.format_time(
                    record.get("cctv_end_time")
                ),
                "elapsed_start": elapsed_start,
            })

        intervals.sort(key=lambda item: item["elapsed_start"])

        # Pair an inactive interval with the following active interval.
        # Not-visible intervals always receive their own row.
        rows = []
        pending_inactive = None

        for interval in intervals:
            state = interval["state"]

            if state == "inactive":
                if pending_inactive is not None:
                    rows.append([
                        pending_inactive["start"],
                        pending_inactive["end"],
                        "",
                        "",
                        "",
                        "",
                    ])

                pending_inactive = interval

            elif state == "active":
                if pending_inactive is not None:
                    rows.append([
                        pending_inactive["start"],
                        pending_inactive["end"],
                        interval["start"],
                        interval["end"],
                        "",
                        "",
                    ])

                    pending_inactive = None
                else:
                    rows.append([
                        "",
                        "",
                        interval["start"],
                        interval["end"],
                        "",
                        "",
                    ])

            elif state == "not_visible":
                if pending_inactive is not None:
                    rows.append([
                        pending_inactive["start"],
                        pending_inactive["end"],
                        "",
                        "",
                        "",
                        "",
                    ])

                    pending_inactive = None

                rows.append([
                    "",
                    "",
                    "",
                    "",
                    interval["start"],
                    interval["end"],
                ])

        if pending_inactive is not None:
            rows.append([
                pending_inactive["start"],
                pending_inactive["end"],
                "",
                "",
                "",
                "",
            ])

        for index, row in enumerate(rows, start=1):
            sheet.append([index, *row])

        widths = [6, 14, 14, 14, 14, 15, 15]

        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                cell.font = Font(size=8)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            sheet.row_dimensions[row[0].row].height = 18

        # Borders around the complete activity table.
        self.apply_borders(
            sheet,
            f"A3:G{max(3, sheet.max_row)}",
        )

        sheet.freeze_panes = "B4"
        sheet.auto_filter.ref = f"A3:G{max(3, sheet.max_row)}"

        self.configure_print(
            sheet,
            f"A1:G{max(3, sheet.max_row)}",
            header_rows="1:3",
        )

    def add_eating_sheet(self, workbook, frequency_records):
        sheet = workbook.create_sheet("Eating Frequency")

        self.style_title(
            sheet,
            "A1:C1",
            "Eating Frequency by Hour",
        )

        headers = ["No.", "Duration", "Eating Frequency"]

        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=3, column=column, value=header)
            self.style_header(cell)

        sheet.row_dimensions[3].height = 25

        for index, record in enumerate(frequency_records, start=1):
            try:
                frequency = int(float(record.get("Frequency") or 0))
            except (TypeError, ValueError):
                frequency = 0

            sheet.append([
                index,
                record.get("Duration", ""),
                frequency,
            ])

        sheet.column_dimensions["A"].width = 9
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 22

        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                cell.font = Font(size=9)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )

            sheet.row_dimensions[row[0].row].height = 20

        # Borders around the complete frequency table.
        self.apply_borders(
            sheet,
            f"A3:C{max(3, sheet.max_row)}",
        )

        sheet.freeze_panes = "A4"
        sheet.auto_filter.ref = f"A3:C{max(3, sheet.max_row)}"

        self.configure_print(
            sheet,
            f"A1:C{max(3, sheet.max_row)}",
            header_rows="1:3",
        )

    def add_visualization_sheet(
        self,
        workbook,
        activity_records,
        eating_records,
        metadata_records,
    ):
        sheet = workbook.create_sheet("Visualization")

        self.style_title(
            sheet,
            "A1:E1",
            "Civet Monitoring Overview",
        )

        metadata = {
            row.get("metadata"): row.get("value")
            for row in metadata_records
        }

        # Recording information
        sheet["A3"] = "Recording Information"
        sheet["A3"].font = Font(bold=True, size=11)

        information = [
            ("Recording Date", metadata.get("recording_date") or "Unavailable"),
            (
                "Recording Start",
                metadata.get("recording_start_timestamp") or "Unavailable",
            ),
            (
                "Recording End",
                metadata.get("recording_end_timestamp") or "Unavailable",
            ),
            (
                "Video Duration",
                self.format_elapsed(
                    self.parse_duration(
                        metadata.get("video_duration_seconds")
                    )
                ),
            ),
        ]

        for row_index, (label, value) in enumerate(information, start=4):
            sheet.cell(row=row_index, column=1, value=label)
            sheet.cell(row=row_index, column=2, value=value)

        # Borders around recording information.
        self.apply_borders(sheet, "A4:B7")

        # Calculate activity durations.
        durations = defaultdict(float)

        for record in activity_records:
            state = record.get("state")

            if state in {"active", "inactive", "not_visible"}:
                durations[state] += self.parse_duration(
                    record.get("duration_seconds")
                )

        active = durations["active"]
        inactive = durations["inactive"]
        not_visible = durations["not_visible"]
        total_observed = active + inactive

        sheet["A10"] = "Activity Duration"
        sheet["A10"].font = Font(bold=True, size=11)

        duration_rows = [
            ("Active", active),
            ("Inactive", inactive),
            ("Not Visible", not_visible),
            ("Observed Time (Active + Inactive)", total_observed),
        ]

        for row_index, (label, seconds) in enumerate(
            duration_rows,
            start=11,
        ):
            sheet.cell(row=row_index, column=1, value=label)
            sheet.cell(
                row=row_index,
                column=2,
                value=self.format_elapsed(seconds),
            )

        # Borders around activity duration information.
        self.apply_borders(sheet, "A11:B14")

        # Eating summary.
        sheet["A17"] = "Confirmed Eating Episodes"
        sheet["A17"].font = Font(bold=True, size=11)
        sheet["B17"] = len(eating_records)

        # Border around eating summary.
        self.apply_borders(sheet, "A17:B17")

        # Chart data: not-visible time is intentionally excluded.
        sheet["D3"] = "Activity"
        sheet["E3"] = "Duration (seconds)"

        sheet["D4"] = "Active"
        sheet["E4"] = active

        sheet["D5"] = "Inactive"
        sheet["E5"] = inactive

        for cell in sheet[3][3:5]:
            self.style_header(cell)

        # Borders around the chart's source data.
        self.apply_borders(sheet, "D3:E5")

        chart = PieChart()
        chart.title = "Active vs Inactive Time"
        chart.style = 10
        chart.height = 7
        chart.width = 11

        data = Reference(
            sheet,
            min_col=5,
            min_row=3,
            max_row=5,
        )

        labels = Reference(
            sheet,
            min_col=4,
            min_row=4,
            max_row=5,
        )

        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)

        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showVal = True

        sheet.add_chart(chart, "D7")

        sheet.column_dimensions["A"].width = 34
        sheet.column_dimensions["B"].width = 24
        sheet.column_dimensions["C"].width = 3
        sheet.column_dimensions["D"].width = 16
        sheet.column_dimensions["E"].width = 20

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center",
                    horizontal="left",
                    wrap_text=True,
                )

                if cell.value is not None and cell.row not in (1, 3, 10, 17):
                    cell.font = Font(size=9)

        for row_index in range(3, 18):
            sheet.row_dimensions[row_index].height = 20

        sheet.row_dimensions[17].height = 25

        self.configure_print(
            sheet,
            "A1:E22",
        )

    def generate(
        self,
        activity_records,
        eating_csv_path,
        frequency_csv_path,
        metadata_csv_path,
    ):
        """Create the final Excel workbook."""

        eating_records = self.read_csv(eating_csv_path)
        frequency_records = self.read_csv(frequency_csv_path)
        metadata_records = self.read_csv(metadata_csv_path)

        workbook = Workbook()

        self.add_activity_sheet(workbook, activity_records)

        self.add_eating_sheet(
            workbook,
            frequency_records,
        )

        self.add_visualization_sheet(
            workbook,
            activity_records,
            eating_records,
            metadata_records,
        )

        workbook.save(self.report_path)

        print(f"Excel report: {self.report_path}")

        return self.report_path
